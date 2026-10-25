"""Build the sprint report (Word + PDF) from the saved results.

Usage: python report/build_report.py   (then the PDF is made with LibreOffice)
"""

import json
import os
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
S = json.load(open(os.path.join(ROOT, "results", "summary.json")))
R = {r["condition"]: r for r in S["summary"]}
SW = {r["random_share"]: r for r in S["sweep"]}
RB = {r["check"]: r for r in S["robustness"]}
SET = S["settings"]
REPO = "https://github.com/emediong-etuk/report-sprint"

TITLE = "One Honest Bot Is Enough: Breaking Up Price Collusion Between Learning Pricing Bots"
AUTHORS = "Emediong Etuk"
AFFIL = "Independent researcher"


def f2(x):
    return f"{x:.2f}"


def ci(row, key):
    m, c = row[key]
    return f"{m:.2f} ± {c:.2f}"


def pval(p):
    return "< 0.001" if p < 0.001 else f"{p:.3f}"


base = R["baseline"]
hedge = R["one_hedge_bot"]
hide = R["hide_prices"]
rnd = R["random_5pct"]
third = R["third_seller"]
N = base["n_seeds"]
drop_hedge = 100 * (1 - hedge["collusion_index"][0] / base["collusion_index"][0])

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1)
    s.top_margin = s.bottom_margin = Inches(0.9)


def para(text, bold=False, italic=False, size=None, align=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def rich(parts, style=None):
    """parts: list of (text, bold) pairs in one paragraph."""
    p = doc.add_paragraph(style=style)
    for text, bold in parts:
        r = p.add_run(text)
        r.bold = bold
    p.paragraph_format.space_after = Pt(6)
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)


def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x0B, 0x0B, 0x0B)


def figure(path, caption, width=6.0):
    doc.add_picture(os.path.join(ROOT, "figures", path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(caption, italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


def table(header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------- Title block ----------------
para(TITLE, bold=True, size=17, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para(AUTHORS, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
para(AFFIL, italic=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("With Apart Research, Poseidon Research and AE Studio. AI Agent Collusion Research Sprint, "
     "October 23-25, 2026. Track 1 (Markets and Collusion), question 1.5.",
     size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

# ---------------- Abstract ----------------
heading("Abstract", 1)
abstract = (
    "Many shops already let computer programs set their prices. Earlier work showed that two simple "
    "learning bots can, with no talking and no instructions, learn to keep prices high together and to "
    "punish a rival that cuts its price. This is a kind of collusion. We ask a simple question: which one "
    "change to the market removes the most collusion at the lowest cost? We rebuild the standard test "
    "market from Calvano et al. (2020) and train two Q-learning bots in it. Then we try four single "
    "changes, one at a time: (1) hide rival prices from the bots, (2) force each bot to pick a random "
    f"price 5% of the time, (3) add a third seller, and (4) swap just one of the two bots for a "
    f"\"no-regret\" bot (the Hedge algorithm) that only tries to do well against whatever prices it sees. "
    f"We run {N} random seeds per setting. Without any change, the bots reach a collusion index of "
    f"{f2(base['collusion_index'][0])} (0 means fair competition, 1 means a full cartel), and "
    f"{round(100 * base['share_punish_cheat'])}% of tested pairs punish a cheater. Swapping one bot for a "
    f"Hedge bot cuts the index to {f2(hedge['collusion_index'][0])}, a {drop_hedge:.0f}% drop, and both "
    "sellers still earn at least fair-competition profit. Forced randomness helps only partly, adding a "
    f"seller helps a little, and hiding rival prices makes collusion worse (index "
    f"{f2(hide['collusion_index'][0])}). Our results suggest that asking one seller to use an "
    "honest, no-regret pricing tool may be a cheap and strong way to break up bot collusion.")
para(abstract)

# ---------------- 1 Introduction ----------------
heading("1. Introduction", 1)
para("Online shops, airlines, hotels and petrol stations use pricing software that changes prices by "
     "itself. When two sellers both use learning software, a worry appears: the programs may learn to "
     "keep prices high together, as if they had agreed to a cartel, even though no human asked them to "
     "and they never send a message. This hurts shoppers. A real study of German petrol stations found "
     "that margins rose when both stations in a local market switched to pricing software "
     "(Assad et al., 2024).")
para("Researchers have shown this effect in a clean computer test (Calvano et al., 2020): two Q-learning "
     "bots, which learn by trial and error, end up charging prices far above the fair, competitive level. "
     "They also learn a \"reward and punishment\" plan: if one bot cuts its price, the other answers with "
     "a price war for a few rounds, and then both go back to high prices. That is exactly how human "
     "cartels keep members in line.")
para("Many fixes have been tested one by one, but rarely side by side in the same market. Our question "
     "(sprint question 1.5) is: which single change removes the most collusion at the least cost? "
     "Our main contributions are:")
bullet(" we compare four single changes in one market, with the same seeds and the same scoring.",
       "One fair test:")
bullet(" we first check that collusion really appears without any change (a positive control), "
       "including the cheat-and-punish test.", "A working baseline:")
bullet(" replacing just one of two bots with a no-regret Hedge bot removes almost all collusion. "
       "Hiding rival prices, a popular idea, makes collusion worse.", "Main finding:")

# ---------------- 2 Related work ----------------
heading("2. Related Work", 1)
para("Calvano et al. (2020) showed that Q-learning pricing bots learn to collude in a repeated market "
     "with logit demand. We copy their market and learning settings. Klein (2021) found similar results "
     "when bots move one after the other. Asker, Fershtman and Pakes (2022) showed that how a bot learns "
     "matters a lot for the prices it ends up with, which is close to our idea that changing the type of "
     "one bot can change the whole market.")
para("Collina et al. (2025) argue that one seller who plays a no-regret strategy can be enough to break "
     "a collusive outcome; we test this idea with learning bots. Keppo et al. (2026) found that mixing in "
     "different kinds of agents lowers collusion between language-model sellers. Xu et al. (2026) report "
     "that a higher sampling temperature weakens links between agents; our \"forced random price\" change "
     "is the simple version of this idea. Fish et al. (2024) showed that language-model pricing agents "
     "can also collude. Hedge itself is a classic no-regret method (Freund and Schapire, 1997).")

# ---------------- 3 Methods ----------------
heading("3. Methods", 1)
heading("3.1 The market", 2)
para("Two sellers (three in one setting) sell similar products. Every round, each seller picks one of "
     "15 prices. Shoppers pick a seller using the standard logit demand model, so the cheaper seller gets "
     "more buyers but not all of them. Each seller earns (price - cost) × buyers. We use exactly the "
     "numbers of Calvano et al. (2020): quality 2, cost 1, outside option 0, and price sensitivity 0.25. "
     f"The fair, competitive price (Nash price) is {SET['benchmarks']['2']['p_nash']:.2f} and the cartel "
     f"price (best joint price) is {SET['benchmarks']['2']['p_mono']:.2f}. The 15 prices are spread evenly "
     "from a bit below the fair price to a bit above the cartel price.")
heading("3.2 The bots", 2)
bullet(" each bot remembers the last prices of all sellers. It learns a score for every price in every "
       f"situation by trial and error (step size {SET['alpha']}, patience {SET['delta']}). At first it "
       f"explores at random a lot; random play fades over time (rate {SET['beta']}). Training lasts "
       f"{SET['T_train']:,} rounds.", "Q-learning bot:")
bullet(" this bot keeps a weight for each price. After every round, it checks how much money each of "
       "its 15 prices would have made against the prices the rival just used, and it shifts weight "
       "toward prices that would have done well. It has no memory of who did what and does not try to "
       f"reward or punish anyone. Learning rate {SET['eta_hedge']:.4f}, the standard textbook choice.",
       "Hedge bot (no-regret):")
heading("3.3 The five settings", 2)
table(["Setting", "What changes", "Who must change"],
      [["No restriction (baseline)", "Two Q-learning bots, nothing changed", "Nobody"],
       ["Hide rival prices", "Bots cannot see past prices", "The platform (all sellers)"],
       ["Force 5% random prices", "Each bot picks a random price 5% of the time, forever", "Every seller's bot"],
       ["Add a third seller", "Three Q-learning bots", "A new seller must enter"],
       ["Swap one bot for Hedge", "One Q-learning bot is replaced by a Hedge bot", "One seller only"]],
      widths=[1.8, 2.9, 1.8])
heading("3.4 How we measure", 2)
para("After training, we watch each market for 10,000 more rounds with learning turned off "
     "(but the forced randomness kept on in that setting). We record:")
bullet(" (profit - fair profit) / (cartel profit - fair profit), averaged over sellers. "
       "0 = fair competition, 1 = full cartel. This is the measure used by Calvano et al.",
       "Collusion index:")
bullet(" shopper welfare scaled the same way: 1 = as good for shoppers as fair competition, "
       "0 = as bad as a full cartel.", "Shopper welfare index:")
bullet(" in the baseline, we let bot A cheat with a one-round price cut and check if bot B answers "
       "with a lower price in the next round. A punishment is a sign of real collusion, not just slow "
       "learning.", "Cheat test:")
bullet(" what the change asks of the market: how many sellers must change and what it costs them.",
       "Cost:")
para(f"Each setting uses the same {N} seeds (0 to {N - 1}). We report the mean and a 95% confidence "
     "interval (± 1.96 standard errors). We compare each setting to the baseline with a two-sided "
     "permutation test (20,000 shuffles). Each run takes about one second on a laptop CPU; the whole "
     "study runs in about 15 minutes with no GPU and no paid API.")

# ---------------- 4 Results ----------------
heading("4. Results", 1)
heading("4.1 Positive control: the bots do collude", 2)
para(f"Without any change, the two Q-learning bots reach a collusion index of {ci(base, 'collusion_index')}. "
     f"Their average price is {f2(base['mean_price'][0])}, well above the fair price of "
     f"{SET['benchmarks']['2']['p_nash']:.2f}. In {round(100 * base['share_runs_index_above_0.5'])}% of runs "
     f"the index is above 0.5. In the cheat test, {round(100 * base['share_punish_cheat'])}% of the "
     f"{base['n_cheat_tests']} tested pairs punish the cheater right away, and prices then climb back up "
     "(Figure 3). This matches Calvano et al. (2020), so our market works as expected.")
heading("4.2 Which single change works best?", 2)
table(["Setting", "Collusion index", "Avg. price", "Shopper welfare", "p vs. baseline", "Who must change"],
      [["No restriction", ci(base, "collusion_index"), f2(base["mean_price"][0]), ci(base, "consumer_index"), "-", "Nobody"],
       ["Hide rival prices", ci(hide, "collusion_index"), f2(hide["mean_price"][0]), ci(hide, "consumer_index"), pval(hide["p_vs_baseline"]), "All (platform)"],
       ["Add a third seller", ci(third, "collusion_index"), f2(third["mean_price"][0]), ci(third, "consumer_index"), pval(third["p_vs_baseline"]), "New entrant"],
       ["Force 5% random prices", ci(rnd, "collusion_index"), f2(rnd["mean_price"][0]), ci(rnd, "consumer_index"), pval(rnd["p_vs_baseline"]), "All sellers"],
       ["Swap one bot for Hedge", ci(hedge, "collusion_index"), f2(hedge["mean_price"][0]), ci(hedge, "consumer_index"), pval(hedge["p_vs_baseline"]), "One seller"]],
      widths=[1.45, 1.1, 0.7, 1.1, 0.8, 1.0])
para("Table 1. Results over " + str(N) + " seeds per setting (mean ± 95% confidence interval). "
     "Note: the third-seller setting is scored against its own three-seller fair and cartel benchmarks.",
     italic=True, size=9.5)
figure("fig1_collusion_by_restriction.png",
       "Figure 1. Collusion index for each single change. Bars show the mean, black lines the 95% "
       "confidence interval, and dots single runs.")
qg = hedge["q_bot_profit_gain"][0]
hg = hedge["hedge_bot_profit_gain"][0]
rich([("Swapping one bot for a Hedge bot works best. ", True),
      (f"The collusion index falls from {f2(base['collusion_index'][0])} to "
       f"{f2(hedge['collusion_index'][0])} ({drop_hedge:.0f}% less), and shopper welfare rises to "
       f"{f2(hedge['consumer_index'][0])}, close to fair competition. Only one seller has to change. "
       f"Both sellers still earn a bit more than fair-competition profit (profit index {f2(hg)} for the "
       f"Hedge seller and {f2(qg)} for the Q-learning seller). The only thing lost is the cartel profit, "
       "which should not exist in the first place. But this also means a seller inside a cartel would not "
       "switch by itself, so the change has to be asked for by a rule or a platform.", False)])
rich([("Forcing random prices helps, but only partly, and it costs every seller. ", True),
      (f"With 5% random prices, the index drops to {f2(rnd['collusion_index'][0])}. Figure 2 shows that "
       f"even 20% random prices still leave an index of {f2(SW[0.2]['collusion_index'][0])}, worse than "
       "one Hedge bot with no randomness at all. Random prices also make the seller's prices jump around "
       "for no reason.", False)])
rich([("Adding a third seller helps only a little ", True),
      (f"(index {f2(third['collusion_index'][0])}). Three bots still learn to keep prices high, as "
       "Calvano et al. also found.", False)])
rich([("Hiding rival prices makes things worse. ", True),
      (f"The index rises to {f2(hide['collusion_index'][0])} (p {pval(hide['p_vs_baseline'])}). Bots that cannot see "
       "the rival can no longer punish cheaters, but they still drift to high prices because each one "
       "learns slowly from a moving target. So taking away information is not a reliable fix, which agrees "
       "with Asker, Fershtman and Pakes (2022).", False)])
figure("fig2_random_price_dose.png",
       "Figure 2. More forced randomness lowers collusion only slowly. Shaded bands show 95% "
       "confidence intervals. The dashed line is one Hedge bot with no forced randomness.", width=5.6)
heading("4.3 Robustness checks for the Hedge result", 2)
rb = RB
table(["Check", "Collusion index", "Shopper welfare"],
      [["Hedge learning rate × 0.1 (slower)", ci(rb["hedge_eta_x0.1"], "collusion_index"), ci(rb["hedge_eta_x0.1"], "consumer_index")],
       ["Hedge learning rate × 10 (faster)", ci(rb["hedge_eta_x10"], "collusion_index"), ci(rb["hedge_eta_x10"], "consumer_index")],
       ["3 sellers, one of them Hedge", ci(rb["three_sellers_1_hedge"], "collusion_index"), ci(rb["three_sellers_1_hedge"], "consumer_index")],
       ["Two Hedge bots (control: no Q-learning)", ci(rb["two_hedge_bots"], "collusion_index"), ci(rb["two_hedge_bots"], "consumer_index")]],
      widths=[3.0, 1.6, 1.6])
para(f"Table 2. Robustness checks ({N} seeds each).", italic=True, size=9.5)
para(f"The Hedge result holds when its learning rate is 10 times slower (index "
     f"{f2(rb['hedge_eta_x0.1']['collusion_index'][0])}) or 10 times faster "
     f"({f2(rb['hedge_eta_x10']['collusion_index'][0])}). In a three-seller market, one Hedge bot cuts the "
     f"index from {f2(third['collusion_index'][0])} to {f2(rb['three_sellers_1_hedge']['collusion_index'][0])}, "
     f"and there the Hedge seller earns the most of all three (profit index "
     f"{f2(rb['three_sellers_1_hedge']['profit_gain_by_seller'][2][0])}). Two Hedge bots together land at "
     f"fair competition ({f2(rb['two_hedge_bots']['collusion_index'][0])}), as theory predicts. This is a "
     "useful negative control: our scoring does not report collusion where none should exist.")
figure("fig3_cheat_test.png",
       "Figure 3. Cheat test in the baseline. Bot A cuts its price once; bot B answers with a lower price "
       "(punishment), then prices climb back toward the cartel level. Gray lines are single runs.",
       width=5.4)

# ---------------- 5 Discussion ----------------
heading("5. Discussion and Conclusion", 1)
para("Our main message is simple: collusion between learning bots needs every bot to play along. "
     "A Q-learning bot learns to keep prices high because the rival rewards high prices and punishes "
     "price cuts. A Hedge bot never does this. It only asks \"which price would have earned the most "
     "against what I just saw?\" So the Q-learning bot finds that high prices are not rewarded, and the "
     "whole market moves back near fair prices.")
para("This points to a cheap policy idea. A regulator or a platform does not need to watch every message "
     "or change every seller's software. It may be enough that some sellers use pricing tools that are "
     "certified as \"no-regret\": tools that try to do well against the market but do not condition on "
     "rewarding or punishing rivals. The seller using such a tool is not hurt by it in our test.")
para("This is also relevant to AI agents built on language models. Those agents may learn or be prompted "
     "into the same reward-and-punish patterns. Our test suggests checking whether one agent with a "
     "simple, history-free decision rule can break a coalition of language-model agents, as Collina et "
     "al. (2025) and Keppo et al. (2026) suggest.")
para("Future work: (1) repeat the test with language-model sellers using the sprint starter code; "
     "(2) add the Hedge seller after collusion has already formed, not from the start; "
     "(3) test bots that know a Hedge seller is present and try to exploit it; and "
     "(4) test larger markets and other demand models.")

# ---------------- Code & data ----------------
heading("Code and Data", 1)
para(f"All code, raw results for every seed, and figures: {REPO}. Run "
     "\"pip install -r requirements.txt\" then \"python src/run_experiment.py\" and "
     "\"python src/make_figures.py\". All code was written during the sprint; no earlier work was used "
     "apart from the published methods we cite.")
heading("LLM Usage Statement", 1)
para("We used an AI coding assistant (Claude Code) to help write the simulation code, run the "
     "experiments, make the figures and draft this report. We checked the numbers in the report against "
     "the saved result files. The authors are responsible for all content.")

# ---------------- References ----------------
heading("References", 1)
refs = [
    "Asker, J., Fershtman, C., and Pakes, A. (2022). Artificial Intelligence, Algorithm Design, and Pricing. AEA Papers and Proceedings, 112, 452-456.",
    "Assad, S., Clark, R., Ershov, D., and Xu, L. (2024). Algorithmic Pricing and Competition: Empirical Evidence from the German Retail Gasoline Market. Journal of Political Economy, 132(3).",
    "Calvano, E., Calzolari, G., Denicolò, V., and Pastorello, S. (2020). Artificial Intelligence, Algorithmic Pricing, and Collusion. American Economic Review, 110(10), 3267-3297.",
    "Collina, N., et al. (2025). Work on no-regret players destabilising collusive equilibria, as cited in the AI Agent Collusion Research Sprint resources (Apart Research, 2026).",
    "Fish, S., Gonczarowski, Y. A., and Shorrer, R. I. (2024). Algorithmic Collusion by Large Language Models. arXiv:2404.00806.",
    "Freund, Y., and Schapire, R. E. (1997). A Decision-Theoretic Generalization of On-Line Learning and an Application to Boosting. Journal of Computer and System Sciences, 55(1), 119-139.",
    "Keppo, J., et al. (2026). Work on mixing agent types in LLM pricing markets, as cited in the AI Agent Collusion Research Sprint resources (Apart Research, 2026).",
    "Klein, T. (2021). Autonomous Algorithmic Collusion: Q-learning under Sequential Pricing. RAND Journal of Economics, 52(3), 538-558.",
    "Xu, et al. (2026). Work on cross-agent dependence and sampling temperature, as cited in the AI Agent Collusion Research Sprint resources (Apart Research, 2026).",
]
for r in refs:
    p = para(r, size=9.5, space_after=3)

# ---------------- Appendix ----------------
heading("Appendix A. Limitations and Dual-Use / Ethical Considerations", 1)
heading("What our result shows and does not show", 2)
bullet(" Our bots are simple Q-learning programs, not language-model agents. The result may not carry "
       "over to AI agents that can talk, plan or read prompts.", "Simple bots only.")
bullet(" One market model (logit demand, 15 prices, 2-3 sellers) with one set of learning settings. "
       "Other markets may behave differently.", "One toy market.")
bullet(" The Hedge seller was present from the start. We did not test adding it after a cartel had "
       "already formed.", "Hedge from the start.")
bullet(" The Q-learning bots did not know a Hedge bot was there. A smarter colluder might try to "
       "exploit or detect it.", "No adaptive colluders.")
bullet(" Our \"cost\" is simple: how many sellers must change and whether they lose profit. We did not "
       "measure real costs such as software, legal or monitoring costs.", "Rough cost measure.")
bullet(f" We stop training after a fixed {SET['T_train']:,} rounds. In the baseline, "
       f"{round(100 * base['stable_share'])}% of runs had stopped changing their plans for the last "
       "500,000 rounds. In the forced-random, three-seller and Hedge settings, the Q-learning bots were "
       "still adjusting, so those results describe behavior after a fixed training time, not a final "
       "settled state.", "Not always settled.")
heading("Dual-use and ethics", 2)
para("Our work studies how to break up collusion, not how to build it. The collusion-forming bots we use "
     "are already fully described in public papers (Calvano et al., 2020), so we add no new way to "
     "collude. One possible misuse: a group of colluding sellers could learn from our results that "
     "Hedge-style rivals break their scheme, and try to detect or push out such rivals. We think the "
     "benefit of telling regulators and platforms about a cheap fix is larger than this risk. No human "
     "data, no real market and no real money were used.")
heading("Appendix B. Full settings", 2)
table(["Setting", "Value"],
      [["Demand", "Logit; quality a = 2, outside option a0 = 0, sensitivity mu = 0.25, cost c = 1"],
       ["Price grid", "15 evenly spaced prices, 10% beyond the fair and cartel prices"],
       ["Q-learning", f"alpha = {SET['alpha']}, delta = {SET['delta']}, exploration = exp(-{SET['beta']} t), memory = last prices of all sellers"],
       ["Training / scoring", f"{SET['T_train']:,} rounds training, {SET['T_eval']:,} rounds scoring"],
       ["Hedge", f"eta = sqrt(8 ln 15 / T) = {SET['eta_hedge']:.4f}, full information"],
       ["Seeds", f"{N} per setting (0-{N - 1}), same seeds in every setting"]],
      widths=[1.6, 4.9])

out_docx = os.path.join(HERE, "report.docx")
doc.save(out_docx)
print("saved", out_docx)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", HERE, out_docx],
               check=True, capture_output=True)
print("saved", os.path.join(HERE, "report.pdf"))
